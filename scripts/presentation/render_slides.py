#!/usr/bin/env python3
"""Render docs/presentation_consortium_45min.html into one 16:9 image per slide, exactly as
the web deck lays it out (1280 x 720 box, normal fonts), and package them as a PDF and a
PowerPoint file with the speaker notes in the notes pane.

Usage:  python3 scripts/presentation/render_slides.py [--scale 2]

A slide whose content is taller than the box is scaled down uniformly (top-left anchored) so
nothing overlaps the footer; short slides keep their natural top-aligned layout. Open Sans is
inlined from a cached copy of the Google Fonts files so the measurement sees the real metrics.
Chromium's --screenshot leaves the footer of the last displayed slide unpainted, so every
chunk is rendered with one blank slide after the real ones. Requires the Playwright
Chromium binary (~/.cache/ms-playwright/chromium-*/chrome-linux*/chrome), Pillow, and
python-pptx (a venv is fine). Outputs (all untracked):
  docs/presentation_consortium_45min.pdf, docs/presentation_consortium_45min.pptx,
  and the per-slide PNGs under the scratch directory given by --out.
"""
from __future__ import annotations
import argparse, glob, importlib.util, json, os, pathlib, re, subprocess, sys, tempfile

ROOT = pathlib.Path(__file__).resolve().parents[2]
HTML = ROOT / "docs" / "presentation_consortium_45min.html"
BUILD = pathlib.Path(__file__).with_name("build_consortium_45min.py")
W, H = 1280, 720

EXPORT_CSS = """
<style id="export">
  html, body { margin:0; padding:0; background:#fff; }
  .deck { max-width:none; margin:0; padding:0; gap:0; display:block; }
  .notes, .notes-bar, footer { display:none !important; }
  .slide { display:none; }
  .slide.export-show { display:flex; flex-direction:column; width:%dpx; height:%dpx; min-height:0; box-sizing:border-box;
           border:0; border-radius:0; margin:0; overflow:hidden; position:relative; }
  .slide.export-show.export-measure { height:auto; overflow:visible; }
  .slide .export-inner { display:flex; flex-direction:column; flex:1 1 auto; min-height:0; transform-origin:top left; }
</style>
""" % (W, H)

EXPORT_JS = """
<script>
(function(){
  const q = new URLSearchParams(location.search);
  const a = parseInt(q.get('from') || '1', 10), b = parseInt(q.get('to') || '9999', 10);
  const slides = Array.from(document.querySelectorAll('.slide'));
  function layout(){
    const dbg = [];
    slides.forEach((s, i) => {
      const n = i + 1; if (n < a || n > b) return;
      s.classList.add('export-show');
      let inner = s.querySelector(':scope > .export-inner');
      if (!inner) {
        inner = document.createElement('div'); inner.className = 'export-inner';
        Array.from(s.children).forEach(c => { if (!c.classList.contains('foot')) inner.appendChild(c); });
        s.insertBefore(inner, s.firstChild);
      }
      inner.style.transform = ''; inner.style.minHeight = '';
      s.classList.add('export-measure');
      const need = inner.getBoundingClientRect().height;   /* natural height, slide unconstrained */
      s.classList.remove('export-measure');
      const cs = getComputedStyle(s);
      const foot = s.querySelector('.foot');
      const avail = s.clientHeight - parseFloat(cs.paddingTop) - (foot ? foot.offsetHeight : 0) - 8;
      let k = 1;
      if (need > avail) { k = avail / need; inner.style.minHeight = need + 'px'; inner.style.transform = 'scale(' + k + ')'; }
      const r = s.getBoundingClientRect();
      dbg.push({n, need: Math.round(need), avail: Math.round(avail), k: +k.toFixed(3), top: Math.round(r.top + scrollY), h: Math.round(r.height),
                foot: foot ? Math.round(foot.getBoundingClientRect().top + scrollY) : null});
    });
    dbg.push({docH: document.documentElement.scrollHeight, fonts: document.fonts ? document.fonts.status : 'n/a'});
    const pass = document.documentElement.hasAttribute('data-export-dbg1') ? 'data-export-dbg2' : 'data-export-dbg1';
    document.documentElement.setAttribute(pass, JSON.stringify(dbg));
  }
  const pad = document.createElement('section'); pad.className = 'slide export-show export-pad';
  slides[slides.length - 1].insertAdjacentElement('afterend', pad);
  layout();
  if (document.fonts && document.fonts.ready) document.fonts.ready.then(layout);
})();
</script>
"""


FONT_CSS_URL = "https://fonts.googleapis.com/css2?family=Open+Sans:wght@400;600;700&display=swap"
UA = "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0 Safari/537.36"


def font_css(out: pathlib.Path) -> str:
    """@font-face rules for Open Sans with the woff2 files inlined as data URIs, so the fonts are
    available at first layout and the render does not depend on Google Fonts answering in time.
    Files are cached under <out>/fonts/; returns "" (keep the Google <link>) if nothing is available."""
    import base64, urllib.request
    fdir = out / "fonts"; fdir.mkdir(parents=True, exist_ok=True)
    css_path = fdir / "opensans.css"
    try:
        if not css_path.exists():
            req = urllib.request.Request(FONT_CSS_URL, headers={"User-Agent": UA})
            css_path.write_bytes(urllib.request.urlopen(req, timeout=30).read())
        css = css_path.read_text()
    except Exception as e:  # noqa: BLE001
        print(f"  warning: no font cache and no download ({e}); using the page's Google Fonts link")
        return ""
    rules, seen = [], {}
    for subset, body in re.findall(r"/\* (\w[\w-]*) \*/\s*@font-face \{(.*?)\}", css, flags=re.S):
        if subset not in ("latin", "latin-ext"):
            continue
        url = re.search(r"url\((https://[^)]+)\)", body).group(1)
        urange = re.search(r"unicode-range: ([^;]+);", body).group(1)
        weight = re.search(r"font-weight: ([^;]+);", body).group(1)
        dst = fdir / f"opensans-{subset}-{pathlib.Path(url).stem[-12:]}.woff2"
        try:
            if not dst.exists():
                urllib.request.urlretrieve(url, dst)
        except Exception as e:  # noqa: BLE001
            print(f"  warning: font download failed ({e}); using the page's Google Fonts link")
            return ""
        key = (subset, dst.name)
        if key in seen:                      # Google serves one variable file for all weights
            seen[key][1].append(weight); continue
        seen[key] = (dst, [weight], urange)
    for (subset, _), (dst, weights, urange) in seen.items():
        ws = sorted({int(w.split()[0]) for w in weights}); wr = f"{ws[0]} {ws[-1]}" if len(ws) > 1 else str(ws[0])
        b64 = base64.b64encode(dst.read_bytes()).decode()
        rules.append(f"@font-face{{font-family:'Open Sans';font-style:normal;font-weight:{wr};font-display:block;"
                     f"src:url(data:font/woff2;base64,{b64}) format('woff2');unicode-range:{urange};}}")
    return "<style id=\"export-fonts\">" + "\n".join(rules) + "</style>"


def chrome() -> str:
    c = sorted(glob.glob(os.path.expanduser("~/.cache/ms-playwright/chromium-*/chrome-linux*/chrome")))
    if not c:
        sys.exit("Chromium not found under ~/.cache/ms-playwright")
    return c[-1]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--scale", type=int, default=2, help="device scale factor (2 = 2560x1440 images)")
    ap.add_argument("--out", default=str(ROOT / "workspace" / "deck_render"))
    a = ap.parse_args()
    out = pathlib.Path(a.out); out.mkdir(parents=True, exist_ok=True)
    html = HTML.read_text()
    n_slides = len(re.findall(r'<section class="slide">', html))
    fonts = font_css(out)
    if fonts:
        html = re.sub(r'<link[^>]*fonts\.g(?:oogleapis|static)\.com[^>]*>\s*', "", html)
        html = html.replace("<style>", fonts + "\n<style>", 1)
    cut = html.rindex("</style>") + len("</style>")          # after the page's last stylesheet
    export = html[:cut] + EXPORT_CSS + html[cut:] + EXPORT_JS
    page = out / "export.html"; page.write_text(export)
    ch = chrome()
    from PIL import Image
    CHUNK = 9
    for start in range(1, n_slides + 1, CHUNK):
        end = min(start + CHUNK - 1, n_slides)
        shot = out / f"chunk-{start:02d}-{end:02d}.png"
        subprocess.run([ch, "--headless=new", "--no-sandbox", "--disable-gpu", "--hide-scrollbars", "--virtual-time-budget=10000",
                        f"--window-size={W},{H * (end - start + 2)}", f"--force-device-scale-factor={a.scale}",
                        f"--screenshot={shot}", f"file://{page}?from={start}&to={end}"],
                       check=True, capture_output=True, timeout=240)
        im = Image.open(shot)
        for k, i in enumerate(range(start, end + 1)):
            im.crop((0, k * H * a.scale, W * a.scale, (k + 1) * H * a.scale)).save(out / f"slide-{i:02d}.png")
        print(f"  slides {start:02d}-{end:02d} rendered ({im.size[0]}x{im.size[1]})")
    # notes + titles from the build module
    spec = importlib.util.spec_from_file_location("b", BUILD); b = importlib.util.module_from_spec(spec); spec.loader.exec_module(b)
    titles = [re.sub(r"<[^>]+>", "", (re.search(r"<h[12]>(.*?)</h[12]>", sl) or [None, ""])[1]) if re.search(r"<h[12]>", sl) else f"Slide {k+1}"
              for k, sl in enumerate(re.findall(r'<section class="slide">(.*?)</section>', html, flags=re.S))]
    assert len(b.NOTES) == n_slides == len(titles)
    # PDF from the images
    from PIL import Image
    ims = [Image.open(out / f"slide-{i:02d}.png").convert("RGB") for i in range(1, n_slides + 1)]
    pdf = ROOT / "docs" / "presentation_consortium_45min.pdf"
    ims[0].save(pdf, save_all=True, append_images=ims[1:], resolution=96 * a.scale)
    print(f"wrote {pdf} ({pdf.stat().st_size} bytes, {n_slides} pages)")
    # PowerPoint with notes
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    for i in range(1, n_slides + 1):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        pic = s.shapes.add_picture(str(out / f"slide-{i:02d}.png"), 0, 0, width=prs.slide_width, height=prs.slide_height)
        pic.name = titles[i - 1]
        s.notes_slide.notes_text_frame.text = b.NOTES[i - 1]
    prs.core_properties.title = "ODELIA Consortium Briefing"; prs.core_properties.author = "JieFu Zhu"
    pptx = ROOT / "docs" / "presentation_consortium_45min.pptx"
    prs.save(pptx); print(f"wrote {pptx} ({pptx.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
